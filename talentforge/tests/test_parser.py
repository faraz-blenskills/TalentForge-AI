"""Whole-document parsing, templates and the score-improvement plan."""
import io
import os
import re
import tempfile
from pathlib import Path

os.environ.setdefault("TALENTFORGE_DB", str(Path(tempfile.mkdtemp()) / "parser-test.db"))
os.environ["PORTAL_AGENT_SCHEDULER"] = "0"
os.environ.pop("ANTHROPIC_API_KEY", None)

import pytest
from fastapi.testclient import TestClient

from app import db
from portal import improve, styles
from portal.doc_extract import extract_docx, extract_pdf, extract_text
from portal.main import app
from portal.resume_parser import classify_heading, normalize_lines
from portal.resume_score import score_resume
from portal.resume_struct import empty_resume, from_text, normalize, parse_with_report, to_text

# Mimics a resume exported with a symbol font: bullets arrive as (cid:127)/DEL, headings in capitals,
# category-labelled skills, a diploma with a polytechnic, and certifications run together on wrapped lines.
GLYPH_RESUME = """ASHA RAO
AI / DATA   SOFTWARE & FULL-STACK   CLOUD / DEVOPS   QA
Pune, India | asha.rao@demo.example | +91 91234 56789 | linkedin.com/in/asha-rao-demo
PROFESSIONAL SUMMARY
Data science undergraduate with hands-on experience in analytics, dashboards and web development. Strong
foundation in Python, SQL and machine learning, with project exposure to Docker and MLflow.
CORE SKILLS
Data & AI: Python, SQL, Pandas, Scikit-learn, Predictive Modeling, Data Analysis
BI & Visualization: Power BI, Tableau, Excel, Dashboard Development
Cloud / DevOps: Docker, AWS S3, MLflow, MLOps fundamentals
WORK EXPERIENCE
Northwind Labs — Data Analytics Intern | Pune | Mar 2026 – May 2026
(cid:127) Analyzed user behaviour across digital platforms to find friction points.
(cid:127) Piloted A/B tests that improved task completion by 15%.
\x7f Documented findings for the product roadmap.
EDUCATION
B.Tech in Artificial Intelligence & Data Science | Lakeview University, Pune | Aug 2024 – May 2027
Current SGPA: 8.65 | Current CGPA: 7.99 | Training Placement Coordinator (TPC)
Diploma in Information Technology | Riverbend Polytechnic, Nashik | Jun 2021 – May 2024
Percentage: 72.56%
PROJECTS
Return Predictor (AI Model) | Python, Scikit-learn, Pandas, Docker, MLflow
(cid:127) Developed a model to identify shipments at risk of return.
(cid:127) Achieved 92% accuracy in predicting delayed returns.
Resume Parser | Mar 2026 – Present
(cid:127) Building an AI-powered resume parsing system.
Data Dashboard | Tableau, SQL
(cid:127) Built interactive dashboards for extensive datasets.
CERTIFICATIONS & TRAINING
AI Dashboards using Microsoft Power BI (cid:127) Python Programming & SQL — Academic Courses (cid:127) AWS Course Certification (cid:127) Advanced Course on
Green Skills and Artificial Intelligence — Skills4Future Program (cid:127) Workshop on Machine Learning — IEEE
Young Professionals (YP) and IEEE Foundation (cid:127) E-Journal Publication — IJSRD
ACHIEVEMENTS
Selected as alternate participant for an international exchange programme sponsored by a government body
LANGUAGES
English, Hindi, Marathi
"""


def parse(text):
    from portal.doc_extract import clean_glyphs
    return parse_with_report(clean_glyphs(text))


# ------------------------------------------------------------------------------ headings
@pytest.mark.parametrize("line,key", [
    ("PROFESSIONAL SUMMARY", "summary"), ("CORE SKILLS", "skills"), ("Technical Skills & Tools", "skills"),
    ("WORK EXPERIENCE", "experience"), ("Internship Experience", "experience"), ("EDUCATION", "education"),
    ("Educational Qualifications", "education"), ("ACADEMIC PROJECTS", "projects"),
    ("CERTIFICATIONS & TRAINING", "certifications"), ("Awards & Achievements", "extras"),
    ("• EDUCATION", "education"), ("Work Experience   3 mos", "experience"), ("Skills:", "skills"),
])
def test_heading_variants_are_recognised(line, key):
    h = classify_heading(line)
    assert h and h[0] == key


@pytest.mark.parametrize("line", [
    "Experience in Python and Flask", "Built a dashboard for the sales team", "Developed projects using Django",
    "• Responsible for projects", "Skills gap analysis for hiring", "B.Tech in Computer Science 2026",
])
def test_body_text_is_not_mistaken_for_a_heading(line):
    assert classify_heading(line) is None


# ------------------------------------------------------------------------------ full parse
def test_symbol_font_resume_is_read_completely():
    r, rep = parse(GLYPH_RESUME)
    b = r["basics"]
    assert b["name"] == "Asha Rao" and b["email"] == "asha.rao@demo.example" and b["phone"].startswith("+91")
    assert b["location"] == "Pune, India" and "SOFTWARE" in b["headline"] and "linkedin.com/in/asha-rao-demo" in b["links"]
    assert len(r["summary"]) > 80
    assert {"Python", "Predictive Modeling", "Dashboard Development", "MLOps fundamentals"} <= set(r["skills"])
    assert len(r["skills"]) >= 14
    assert len(r["experience"]) == 1
    e = r["experience"][0]
    assert "Intern" in e["role"] and "Northwind" in e["company"] and e["start"].startswith("Mar") and e["end"].startswith("May")
    assert len(e["bullets"]) == 3 and not any("cid" in x for x in e["bullets"])
    assert len(r["education"]) == 2
    assert "Polytechnic" in r["education"][1]["institution"]
    assert "Aug 2024" in r["education"][0]["year"] and "SGPA" in r["education"][0]["details"]
    assert [p["name"] for p in r["projects"]][0].startswith("Return Predictor (AI Model)")   # bracket kept
    assert len(r["projects"]) == 3 and sum(len(p["bullets"]) for p in r["projects"]) == 4
    assert len(r["certifications"]) == 6
    assert any("Green Skills" in c and c.startswith("Advanced Course on") for c in r["certifications"])
    assert any("IEEE Foundation" in c for c in r["certifications"])
    titles = [x["title"] for x in r["extras"]]
    assert "Achievements" in titles and "Languages" in titles
    assert rep["counts"]["lines_read"] > 25 and not rep["warnings"]


def test_profile_builder_export_layout():
    text = """Jane Doe
Data Analyst / AI Intern
Pune, India
•
jane@demo.example   +91 90000 11111
•
SUMMARY
Motivated analyst with hands-on experience in dashboards and machine learning.
•
Work Experience   3 mos
Analyst Intern at Contoso Ltd, Pune
Mar 2026 - May 2026   3mos
- Analyzed funnel data to find drop-offs.
- Built weekly dashboards for managers.
• EDUCATION
B.Tech in Data Science   Aug 2024 - May 2027
Greenfield University   Pune
•
TECHNICAL SKILLS
Programming
Python   Java   SQL
Data Analytics Tools
Tableau   Excel   Power BI
"""
    r, rep = parse(text)
    assert r["basics"]["name"] == "Jane Doe" and r["basics"]["location"] == "Pune, India"
    assert len(r["experience"]) == 1 and r["experience"][0]["company"].startswith("Contoso")
    assert len(r["experience"][0]["bullets"]) == 2
    assert len(r["education"]) == 1 and r["education"][0]["institution"].startswith("Greenfield")
    assert {"Python", "Java", "SQL", "Tableau", "Excel", "Power BI"} <= set(r["skills"])
    assert "Programming" not in r["skills"] and "Data Analytics Tools" not in r["skills"]   # labels, not skills
    assert rep["counts"]["lines_read"] >= 10


def test_numeric_dates_and_comma_separated_header():
    text = """Rohit Sen
rohit@demo.example | 9876543210
EXPERIENCE
Software Engineer, Acme Pvt Ltd, 06/2023 - Present
- Built REST APIs serving 3 internal teams.
Junior Developer | Beta Systems | 2021 - 2023
- Maintained a billing module.
EDUCATION
B.E. Computer Engineering, Test College of Engineering, 2021
"""
    r, _ = parse(text)
    assert [e["role"] for e in r["experience"]] == ["Software Engineer", "Junior Developer"]
    assert r["experience"][0]["company"].startswith("Acme") and r["experience"][0]["end"] == "Present"
    assert r["experience"][1]["start"] == "2021" and r["experience"][1]["end"] == "2023"
    assert r["education"][0]["year"] == "2021" and "College" in r["education"][0]["institution"]


def test_wrapped_bullets_are_rejoined():
    text = """Neha Iyer
neha@demo.example
EXPERIENCE
Analyst | Delta Corp | Jan 2024 - Dec 2024
- Designed and implemented an automated reporting pipeline that reduced manual effort for the finance
  team and improved accuracy of month-end numbers across five regions.
- Trained 12 new joiners.
"""
    r, _ = parse(text)
    assert len(r["experience"]) == 1
    bullets = r["experience"][0]["bullets"]
    assert len(bullets) == 2 and "five regions" in bullets[0]


def test_nothing_is_dropped_when_there_are_no_headings():
    text = "Kavya Menon\nkavya@demo.example\nI built things with Python.\nI also led a student club of 40 members.\n" \
           "Worked at Acme for two years on payments.\n"
    r, rep = parse(text)
    assert any(ex["title"] == "Resume content" and len(ex["items"]) >= 2 for ex in r["extras"])
    assert rep["warnings"]
    assert "student club" in to_text(r)


def test_unparsed_sections_survive_as_extras():
    text = GLYPH_RESUME + "\nHOBBIES\nChess, trekking, open-source\nPUBLICATIONS\nE-Journal paper on resume parsing\n"
    r, _ = parse(text)
    assert {"Hobbies", "Publications"} <= {x["title"] for x in r["extras"]}


# ------------------------------------------------------------------------------ extraction
def test_docx_tables_headers_and_lists_are_read():
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Maya Pillai | maya@demo.example | +91 90000 22222"
    d.add_paragraph("SKILLS")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Languages"
    t.rows[0].cells[1].text = "Python, Java, SQL"
    d.add_paragraph("EXPERIENCE")
    d.add_paragraph("Engineer | Omega Ltd | 2022 - 2024")
    d.add_paragraph("Built a data pipeline", style="List Bullet")
    d.add_paragraph("EDUCATION")
    d.add_paragraph("B.Tech in IT, Test University, 2022")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_docx(buf.getvalue())
    assert "Maya Pillai" in text and "Python, Java, SQL" in text and "Built a data pipeline" in text
    r = from_text(text)
    assert r["basics"]["email"] == "maya@demo.example"
    assert r["experience"] and r["experience"][0]["bullets"] == ["Built a data pipeline"]
    assert {"Python", "Java", "SQL"} <= set(r["skills"])


def test_two_column_pdf_is_read_column_by_column():
    matplotlib = pytest.importorskip("matplotlib")
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    matplotlib.rcParams["pdf.fonttype"] = 42
    fig = plt.figure(figsize=(8.27, 11.69))
    left = ["SKILLS", "Python", "SQL", "Docker", "Git", "Linux", "Excel", "EDUCATION", "B.Tech in IT", "Test University",
            "2022 - 2026", "Pune", "CGPA 8.5", "Placement TPC"]
    right = ["EXPERIENCE", "Engineer | Omega Ltd | 2022 - 2024", "- Built data pipelines", "- Wrote unit tests",
             "- Mentored interns", "- Cut costs by 20%", "PROJECTS", "Chat Bot | Python", "- Built an FAQ bot",
             "- Deployed with Docker", "- Served 200 users", "- Added logging"]
    for i, ln in enumerate(left):
        fig.text(0.06, 0.93 - i * 0.045, ln, fontsize=11)
    for i, ln in enumerate(right):
        fig.text(0.52, 0.93 - i * 0.045, ln, fontsize=11)
    buf = io.BytesIO()
    fig.savefig(buf, format="pdf")
    plt.close(fig)
    text = extract_pdf(buf.getvalue())
    # the whole left column must come before the right column (not interleaved line by line)
    assert text.index("Placement TPC") < text.index("EXPERIENCE")
    r = from_text(text)
    assert r["experience"] and len(r["projects"]) == 1 and r["education"]


def test_bullet_glyphs_are_normalised():
    t = extract_text("(cid:127) one\n two\n● three".encode(), "x.txt")
    assert t.count("•") == 3 and "cid" not in t


# ------------------------------------------------------------------------------ templates + plan
def test_three_templates_with_distinct_fonts():
    cat = styles.catalogue()
    assert len(cat) == 3 and len({t["font"] for t in cat}) == 3
    assert styles.resolve("nope", "red") == (styles.DEFAULT_TEMPLATE, styles.TEMPLATES[styles.DEFAULT_TEMPLATE]["default_accent"])
    assert styles.resolve("classic", "#1F3A5F") == ("classic", "#1f3a5f")


def test_docx_follows_template_font():
    import docx
    from portal.export import resume_to_docx
    r = normalize(parse(GLYPH_RESUME)[0])
    for tid, font in (("classic", "Georgia"), ("modern", "Calibri"), ("executive", "Arial")):
        d = docx.Document(io.BytesIO(resume_to_docx(r, tid, "#1f3a5f")))
        assert d.styles["Normal"].font.name == font
        txt = "\n".join(p.text for p in d.paragraphs)
        assert "Achievements".upper() in txt.upper() and "Return Predictor" in txt


def test_fresher_with_three_deep_projects_can_reach_full_marks():
    r = empty_resume("Asha Rao", "asha@demo.example")
    r["basics"].update(phone="+91 90000 00000", location="Pune, India", headline="Data Analyst")
    r["summary"] = ("Data science graduate with hands-on experience in Python, SQL and dashboards, backed by "
                    "three shipped projects that demonstrate applied, results-oriented analytics work.")
    r["skills"] = ["Python", "SQL", "Pandas", "Tableau", "Excel", "Git", "Docker", "Power BI"]
    r["education"] = [{"degree": "B.Tech in Data Science", "institution": "Test University", "year": "2026", "details": ""}]
    r["projects"] = [{"name": f"Project {i}", "description": "Python, SQL", "bullets": [
        f"Built a pipeline that processed {10 * (i + 1)}k records and cut manual work by 40%",
        "Created dashboards used by 3 teams to track weekly KPIs"]} for i in range(3)]
    s = score_resume(r)
    assert s["score"] >= 90, s["checks"]


def test_improvement_plan_is_measured_and_apply_never_lowers_score():
    r = parse(GLYPH_RESUME.replace("Documented findings", "Responsible for documenting findings"))[0]
    r["basics"]["phone"] = ""
    plan = improve.build_plan(r)
    kinds = {i["id"].split(":")[0]: i for i in plan["items"]}
    assert "contact" in kinds and kinds["contact"]["gain"] > 0
    assert "verbs" in kinds and kinds["verbs"]["changes"]
    assert plan["projected_all"] >= plan["projected_auto"] >= plan["score"]
    assert plan["projected_all"] >= 95
    accepted = {i["id"] for i in plan["items"] if i["kind"] != "manual"}
    inputs = {"contact:phone": "+91 98765 00000", "summary": "Data analyst with Python and SQL experience across "
              "three projects and one internship, focused on dashboards and machine learning for decisions."}
    new = improve.apply_plan(r, plan, accepted, inputs)
    assert new["basics"]["phone"] == "+91 98765 00000"
    assert score_resume(new)["score"] > plan["score"]
    assert r["basics"]["phone"] == ""        # original not mutated


# ------------------------------------------------------------------------------ routes
def _register(c, email):
    r = c.post("/register", data={"name": "Asha Rao", "email": email, "password": "longenough1", "consent": "1"},
               follow_redirects=False)
    assert r.status_code == 303
    return db.one("SELECT id FROM users WHERE email=?", (email,))["id"]


def test_upload_shows_parse_report_and_improvement_plan():
    with TestClient(app) as c:
        _register(c, "parser-route@demo.example")
        r = c.post("/resumes/upload", files={"resume": ("asha.txt", GLYPH_RESUME.replace("(cid:127)", "•").encode(),
                                                          "text/plain")},
                   data={"template": "executive", "accent": "#8a1c3d"}, follow_redirects=False)
        assert r.status_code == 303
        rid = int(r.headers["location"].rsplit("/", 1)[1])
        page = c.get(f"/resumes/{rid}")
        assert page.status_code == 200
        assert "What we read from your file" in page.text and "Show the original text we extracted" in page.text
        assert "t-executive" in page.text and "--rc: #8a1c3d" in page.text
        assert "Return Predictor" in page.text and "E-Journal Publication" in page.text     # nothing dropped
        assert "Path to 99%" in page.text

        improve_page = c.get(f"/resumes/{rid}/improve")
        assert improve_page.status_code == 200 and "Your path to 99%" in improve_page.text

        # restyle, then apply the plan as a copy
        assert c.post(f"/resumes/{rid}/style", data={"template": "classic", "accent": "#1f3a5f"},
                      follow_redirects=False).status_code == 303
        assert "t-classic" in c.get(f"/resumes/{rid}").text
        plan = improve.build_plan(parse(GLYPH_RESUME)[0])
        accept = [i["id"] for i in plan["items"] if i["kind"] == "auto"] or ["summary"]
        r2 = c.post(f"/resumes/{rid}/improve", data={"accept": accept, "copy": "1"}, follow_redirects=False)
        assert r2.status_code == 303
        new_id = int(r2.headers["location"].split("/")[2])
        assert new_id != rid


def test_builder_creation_with_template_choice():
    with TestClient(app) as c:
        _register(c, "builder-route@demo.example")
        page = c.get("/resumes")
        assert page.status_code == 200 and page.text.count('name="template"') >= 3
        for tid in ("classic", "modern", "executive"):
            r = c.post("/resumes/new", data={"title": f"{tid} cv", "template": tid, "accent": "#0f766e"},
                       follow_redirects=False)
            rid = int(r.headers["location"].split("/")[2])
            edit = c.get(f"/resumes/{rid}/edit")
            assert edit.status_code == 200 and f'value="{tid}" selected' in edit.text.replace("\n", " ") \
                or f'<option value="{tid}" selected>' in edit.text
            saved = c.post(f"/api/resumes/{rid}", json={"data": empty_resume("A B", "a@b.co"), "template": "classic",
                                                         "accent": "#7a1f2b"})
            assert saved.status_code == 200
            assert "t-classic" in c.get(f"/resumes/{rid}").text
