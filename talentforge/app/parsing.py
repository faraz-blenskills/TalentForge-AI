"""Resume text extraction and structured parsing (rule-based, deterministic).

Output is a plain dict so it can be stored as JSON, edited by the user and re-scored.
"""
from __future__ import annotations

import io
import re
from datetime import date

from .taxonomy import build_alias_index

SECTION_HEADS = {
    "skills": r"(technical\s+)?skills?|core competencies|technologies|tech stack",
    "education": r"education|academic(s| background| qualifications?)?|qualifications?",
    "experience": r"(work |professional )?experience|employment|internships?|work history",
    "projects": r"(academic |personal |key )?projects?",
    "certifications": r"certifications?|certificates?|courses|licen[cs]es?",
    "summary": r"summary|objective|profile|about me",
    "other": r"achievements?|awards?|interests?|hobbies|languages|references|publications?",
}
_HEAD_RE = {
    k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.I) for k, v in SECTION_HEADS.items()
}

DEGREES = [
    (4, r"ph\.?\s?d|doctorate"),
    (3, r"m\.?\s?tech|m\.?\s?e\b|m\.?\s?sc|m\.?\s?c\.?\s?a\b|\bmca\b|\bmba\b|master(?:'s|s)?\b|post\s?graduate"),
    (2, r"b\.?\s?tech|b\.?\s?e\b|b\.?\s?sc|\bbca\b|\bb\.?\s?com\b|bachelor(?:'s|s)?\b|undergraduate|b\.?\s?e\.?\s?\("),
    (1, r"diploma|polytechnic"),
]
CS_FIELDS = re.compile(
    r"computer|information technology|\bit\b|software|data science|artificial intelligence|"
    r"\bai\b|machine learning|cyber|electronics|electrical|\bcse\b|\bise\b|statistics|mathematics",
    re.I,
)
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
RANGE_RE = re.compile(
    r"(?:(?P<m1>[A-Za-z]{3,9})\.?,?\s+)?(?P<y1>(?:19|20)\d{2})\s*(?:-|–|—|to)\s*"
    r"(?:(?P<m2>[A-Za-z]{3,9})\.?,?\s+)?(?P<y2>(?:19|20)\d{2}|present|current|now|ongoing)",
    re.I,
)
EXPLICIT_YEARS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b", re.I)
BULLET_RE = re.compile(r"^\s*(?:[-•*▪●‣◦·]|\d+[.)])\s*")


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return "\n".join(lines)
    return data.decode("utf-8", errors="ignore")


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for ln in lines:
        stripped = ln.strip()
        matched = None
        if 0 < len(stripped) <= 40:
            for key, rx in _HEAD_RE.items():
                if rx.match(stripped):
                    matched = key
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(ln)
    return sections


def _find_skills(text: str) -> list[dict]:
    low = text.lower()
    lines = text.splitlines()
    found: dict[str, dict] = {}
    consumed: list[tuple[int, int]] = []
    for alias, canon in build_alias_index():
        pat = re.compile(r"(?<![a-z0-9+#.])" + re.escape(alias) + r"(?![a-z0-9+#])", re.I)
        for m in pat.finditer(low):
            if any(s <= m.start() < e for s, e in consumed):
                continue
            consumed.append((m.start(), m.end()))
            if canon not in found:
                line_no = low.count("\n", 0, m.start())
                ev = lines[line_no].strip() if line_no < len(lines) else alias
                found[canon] = {"name": canon, "evidence": ev[:160], "mentions": 1}
            else:
                found[canon]["mentions"] += 1
    return sorted(found.values(), key=lambda s: (-s["mentions"], s["name"]))


def _parse_education(lines: list[str]) -> tuple[list[dict], int, bool]:
    entries, best_level, cs_related = [], 0, False
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        level = 0
        for lvl, rx in DEGREES:
            if re.search(rx, s, re.I):
                level = max(level, lvl)
        if level:
            year = re.findall(r"(?:19|20)\d{2}", s)
            field_cs = bool(CS_FIELDS.search(s))
            entries.append({"text": s[:160], "level": level, "year": year[-1] if year else "",
                            "cs_related": field_cs})
            best_level = max(best_level, level)
            cs_related = cs_related or field_cs
    # a degree line without a field often has the field on the next line
    if best_level and not cs_related:
        blob = " ".join(lines[:6])
        cs_related = bool(CS_FIELDS.search(blob))
    return entries, best_level, cs_related


def _months_between(y1, m1, y2, m2) -> float:
    return max((y2 - y1) * 12 + (m2 - m1), 1)


def _parse_experience(lines: list[str], full_text: str) -> tuple[float, list[str]]:
    today = date.today()
    entries: list[str] = []
    months_total = 0.0
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        m = RANGE_RE.search(s)
        if m:
            y1 = int(m.group("y1"))
            m1 = MONTHS.get((m.group("m1") or "")[:3].lower(), 1)
            end = m.group("y2").lower()
            if end.isdigit():
                y2 = int(end)
                m2 = MONTHS.get((m.group("m2") or "")[:3].lower(), 12 if y2 > y1 else 1)
            else:
                y2, m2 = today.year, today.month
            months_total += _months_between(y1, m1, y2, m2)
            entries.append(s[:160])
        elif re.search(r"\b(intern|engineer|developer|analyst|trainee|associate|consultant)\b", s, re.I) \
                and not BULLET_RE.match(s):
            entries.append(s[:160])
    years = round(months_total / 12, 1)
    if years == 0:
        explicit = [float(x) for x in EXPLICIT_YEARS_RE.findall(full_text)]
        plausible = [x for x in explicit if x <= 30]
        if plausible:
            years = max(plausible)
        elif any(re.search(r"intern", e, re.I) for e in entries):
            years = 0.5
    return min(years, 30.0), entries[:8]


def _parse_list(lines: list[str], min_len: int = 6) -> list[str]:
    out = []
    for ln in lines:
        s = BULLET_RE.sub("", ln).strip()
        if len(s) >= min_len:
            out.append(s[:200])
    return out


def _parse_projects(lines: list[str]) -> list[str]:
    heads, bullets = [], []
    for ln in lines:
        s = ln.rstrip()
        if not s.strip():
            continue
        if BULLET_RE.match(s):
            bullets.append(BULLET_RE.sub("", s).strip()[:200])
        elif len(s.strip()) >= 6:
            heads.append(s.strip()[:200])
    return heads if heads else bullets


def parse_resume(text: str) -> dict:
    lines = [ln for ln in text.splitlines()]
    sections = _split_sections(lines)
    warnings: list[str] = []

    header = [l.strip() for l in sections.get("header", []) if l.strip()]
    name = ""
    for l in header[:4]:
        if EMAIL_RE.search(l) or re.search(r"\d{4,}", l):
            continue
        words = l.split()
        if 1 < len(words) <= 4 and all(re.match(r"^[A-Za-z.'\-]+$", w) for w in words):
            name = l.title() if l.isupper() else l
            break
    email_m = EMAIL_RE.search(text)
    email = email_m.group(0) if email_m else ""

    skills = _find_skills(text)
    if "skills" not in sections:
        warnings.append("No dedicated Skills section found; skills were detected from the full text.")
    if not skills:
        warnings.append("No known skills detected. Please add them manually.")

    edu_entries, edu_level, cs_related = _parse_education(
        sections.get("education", []) or lines)
    if not edu_entries:
        warnings.append("Could not detect an education entry.")

    exp_lines = sections.get("experience", [])
    years, exp_entries = _parse_experience(exp_lines if exp_lines else [], text)
    if not exp_lines and years == 0:
        warnings.append("No experience section detected (fine for freshers).")

    projects = _parse_projects(sections.get("projects", []))
    if not projects:
        warnings.append("No projects detected. Projects strongly improve employability for freshers.")
    certs = _parse_list(sections.get("certifications", []))

    return {
        "name": name,
        "email": email,
        "skills": skills,
        "education": edu_entries,
        "highest_degree_level": edu_level,
        "cs_related": cs_related,
        "experience_years": years,
        "experience_entries": exp_entries,
        "projects": projects,
        "certifications": certs,
        "warnings": warnings,
    }
